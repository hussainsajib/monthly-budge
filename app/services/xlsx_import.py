"""One-off import of the legacy Excel 'Transactions' sheet."""

from __future__ import annotations

import hashlib
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.money import to_cents
from app.models import EXPENSE, Account, Category, Transaction

SHEET = "Transactions"
UNCATEGORIZED_GROUP = "Uncategorized"


def _get_or_create_category(db: Session, name: str, cache: dict[str, Category]) -> Category:
    """Find a category by name; unknown names go under an 'Uncategorized' section."""
    if name in cache:
        return cache[name]
    section = db.scalar(select(Category).where(Category.name == UNCATEGORIZED_GROUP, Category.parent_id.is_(None)))
    if section is None:
        section = Category(name=UNCATEGORIZED_GROUP, kind=EXPENSE, sort_order=900)
        db.add(section)
        db.flush()
    cat = Category(name=name, parent_id=section.id, kind=EXPENSE)
    db.add(cat)
    db.flush()
    cache[name] = cat
    return cat


def import_workbook(db: Session, path: Path) -> tuple[int, int]:
    """Import rows; returns (created, skipped_existing). Safe to run repeatedly."""
    wb = load_workbook(path, data_only=False, read_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook has no '{SHEET}' sheet")

    categories = {c.name: c for c in db.scalars(select(Category).where(Category.parent_id.is_not(None)))}
    accounts = {a.name: a for a in db.scalars(select(Account))}
    existing_hashes = set(db.scalars(select(Transaction.import_hash).where(Transaction.import_hash.is_not(None))))

    seen: Counter[str] = Counter()
    created = skipped = 0
    for row in wb[SHEET].iter_rows(min_row=2, values_only=True):
        when, description, category_name, amount, account_name, notes = (list(row) + [None] * 6)[:6]
        if when is None or not description or not category_name or amount in (None, ""):
            continue
        if isinstance(when, datetime):
            when = when.date()
        if not isinstance(when, date):
            continue

        cents = to_cents(amount if isinstance(amount, str) else str(amount))
        base = f"xlsx|{when}|{str(description).strip().lower()}|{category_name}|{cents}"
        digest = hashlib.sha1(f"{base}|{seen[base]}".encode()).hexdigest()
        seen[base] += 1
        if digest in existing_hashes:
            skipped += 1
            continue

        account = accounts.get(str(account_name).strip()) if account_name else None
        db.add(
            Transaction(
                date=when,
                description=str(description).strip()[:200],
                category_id=_get_or_create_category(db, str(category_name).strip(), categories).id,
                account_id=account.id if account else None,
                amount_cents=cents,
                notes=(str(notes).strip() or None) if notes else None,
                import_hash=digest,
            )
        )
        created += 1
    db.commit()
    return created, skipped
